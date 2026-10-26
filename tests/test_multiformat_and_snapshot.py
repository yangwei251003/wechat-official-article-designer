from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from article_schema import scaffold_from_source  # noqa: E402
from compile_wechat import compile_article  # noqa: E402
from parse_inputs import parse_file  # noqa: E402


class MultiformatIntegrationTests(unittest.TestCase):
    def test_docx_extracts_paragraphs_tables_and_image_inventory(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "article.docx"
            document = Document()
            document.add_heading("Word标题", level=1)
            document.add_paragraph("Word正文。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "数值"
            table.cell(1, 0).text = "留存"
            table.cell(1, 1).text = "34%"
            document.save(path)
            result = parse_file(path, root=path.parent)
            self.assertEqual("parsed", result["status"])
            self.assertIn("Word正文。", result["text"])
            self.assertEqual("34%", result["tables"][0][1][1])

    def test_xlsx_preserves_formula_as_source_data(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "data.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "数据"
            sheet.append(["基数", "增长", "结果"])
            sheet.append([100, 0.2, "=A2*(1+B2)"])
            workbook.save(path)
            result = parse_file(path, root=path.parent)
            self.assertEqual("=A2*(1+B2)", result["tables"][0]["rows"][1][2])

    def test_scanned_like_pdf_without_text_reports_degradation(self) -> None:
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=400, height=400)
            with path.open("wb") as handle:
                writer.write(handle)
            result = parse_file(path, root=path.parent)
            self.assertEqual("degraded", result["status"])
            self.assertTrue(any("文本层" in warning for warning in result["warnings"]))

    def test_image_metadata_has_dimensions_frames_and_color(self) -> None:
        from PIL import Image

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "image.png"
            Image.new("RGB", (640, 360), (49, 95, 88)).save(path)
            result = parse_file(path, root=path.parent)
            self.assertEqual((640, 360), (result["width"], result["height"]))
            self.assertEqual(1, result["frames"])
            self.assertEqual("#315f58", result["dominant_color"])


class SnapshotTests(unittest.TestCase):
    def test_strict_wechat_compiler_snapshot(self) -> None:
        article = scaffold_from_source("# 快照\n\n正文。\n", title="快照")
        output = compile_article(article)["html"]
        digest = hashlib.sha256(output.encode("utf-8")).hexdigest()
        self.assertEqual("1a0b1e9c13acdc51a9db16ff8cb5a1bd8630384a6c8f111e713457edd83cf3c1", digest)


if __name__ == "__main__":
    unittest.main(verbosity=2)
