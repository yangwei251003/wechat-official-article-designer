from __future__ import annotations

import argparse
import csv
import html
import json
import mimetypes
import re
import shutil
import subprocess
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from common import now_iso, read_text, safe_resolve, sha256_file, write_json


class SafeTextHTMLParser(HTMLParser):
    BLOCKED = {"script", "style", "noscript", "iframe", "object", "embed"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in self.BLOCKED:
            self.depth += 1
        elif not self.depth and tag.lower() in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "br", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self.BLOCKED and self.depth:
            self.depth -= 1
        elif not self.depth and tag.lower() in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.depth:
            self.parts.append(data)


def _base(path: Path) -> dict[str, Any]:
    return {"schema_version": "1.0", "source_file": path.name, "sha256": sha256_file(path), "size": path.stat().st_size, "media_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream", "parsed_at": now_iso(), "text": "", "headings": [], "tables": [], "images": [], "warnings": [], "status": "parsed", "untrusted_external_content": False}


def _parse_html(path: Path, result: dict[str, Any]) -> None:
    parser = SafeTextHTMLParser()
    parser.feed(read_text(path))
    result["text"] = re.sub(r"\n{3,}", "\n\n", html.unescape("".join(parser.parts))).strip()
    result["untrusted_external_content"] = True
    result["warnings"].append("HTML/网页文本仅作为不可信数据读取；脚本与嵌入对象未执行")


def _parse_docx(path: Path, result: dict[str, Any]) -> None:
    try:
        from docx import Document
    except ImportError:
        result.update(status="degraded", warnings=["缺少python-docx；请另存为Markdown或PDF"])
        return
    document = Document(path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    result["text"] = "\n\n".join(paragraphs)
    result["headings"] = [{"level": int(p.style.name[-1]) if p.style.name[-1:].isdigit() else 1, "text": p.text.strip()} for p in document.paragraphs if p.text.strip() and p.style and p.style.name.lower().startswith("heading")]
    result["tables"] = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    result["images"] = [{"relationship_id": rel_id, "filename": rel.target_ref} for rel_id, rel in document.part.rels.items() if "image" in rel.reltype]


def _parse_pdf(path: Path, result: dict[str, Any]) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:
        result.update(status="degraded", warnings=["缺少pypdf；请提供Markdown或页面截图"])
        return
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            result.update(status="blocked", warnings=["PDF受密码保护，未尝试绕过"])
            return
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        result["text"] = "\n\n".join(text for text in pages if text)
        result["page_count"] = len(reader.pages)
        if not result["text"]:
            result.update(status="degraded", warnings=["PDF没有可提取文本层；请提供页面截图进行视觉/OCR分析"])
    except Exception as exc:
        result.update(status="failed", warnings=[f"PDF解析失败：{type(exc).__name__}"])


def _parse_csv(path: Path, result: dict[str, Any]) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [row for row in csv.reader(handle)]
    result["tables"] = [rows]
    result["text"] = "\n".join(" | ".join(row) for row in rows)


def _parse_xlsx(path: Path, result: dict[str, Any]) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        result.update(status="degraded", warnings=["缺少openpyxl；请另存为CSV"])
        return
    workbook = load_workbook(path, read_only=True, data_only=False)
    tables = []
    try:
        for sheet in workbook.worksheets:
            rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
            tables.append({"sheet": sheet.title, "rows": rows})
    finally:
        workbook.close()
    result["tables"] = tables
    result["text"] = "\n\n".join(f"[{table['sheet']}]\n" + "\n".join(" | ".join("" if value is None else str(value) for value in row) for row in table["rows"]) for table in tables)


def _parse_image(path: Path, result: dict[str, Any]) -> None:
    try:
        from PIL import Image, ImageStat
    except ImportError:
        result.update(status="degraded", warnings=["缺少Pillow，无法分析图片尺寸和颜色"])
        return
    with Image.open(path) as image:
        rgb = image.convert("RGB").resize((1, 1))
        color = tuple(ImageStat.Stat(rgb).mean)
        result.update(width=image.width, height=image.height, frames=getattr(image, "n_frames", 1), dominant_color="#%02x%02x%02x" % tuple(int(x) for x in color))
    result["warnings"].append("未配置本地OCR；已保留图片元数据，可由视觉模型或授权OCR继续处理")


def _parse_media(path: Path, result: dict[str, Any]) -> None:
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        result.update(status="degraded", warnings=["未找到ffprobe；无法自动提取音视频时长、编码和转录，请提供文字稿或安装ffmpeg"])
        return
    completed = subprocess.run([ffprobe, "-v", "error", "-show_entries", "format=duration:stream=codec_type,codec_name", "-of", "json", str(path)], capture_output=True, text=True, timeout=30)
    if completed.returncode:
        result.update(status="failed", warnings=["ffprobe无法读取媒体文件"])
    else:
        result["media_info"] = json.loads(completed.stdout)
        result["warnings"].append("已提取媒体元数据；转录需要本地模型或用户授权的外部服务")


def parse_file(path: Path, *, root: Path | None = None) -> dict[str, Any]:
    root = (root or path.parent).resolve()
    path = safe_resolve(path, root)
    result = _base(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        result["text"] = read_text(path)
    elif suffix in {".html", ".htm"}:
        _parse_html(path, result)
    elif suffix == ".docx":
        _parse_docx(path, result)
    elif suffix == ".pdf":
        _parse_pdf(path, result)
    elif suffix == ".csv":
        _parse_csv(path, result)
    elif suffix == ".xlsx":
        _parse_xlsx(path, result)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        _parse_image(path, result)
    elif suffix == ".svg":
        result["text"] = read_text(path)
        result["warnings"].append("SVG在使用前必须运行sanitize_svg.py并输出静态PNG降级")
    elif suffix in {".mp3", ".wav", ".m4a", ".aac", ".mp4", ".mov", ".webm"}:
        _parse_media(path, result)
    else:
        result.update(status="unsupported", warnings=[f"暂不支持{suffix or '无扩展名'}；请转换为支持格式"])
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="把文档、表格、图片和媒体解析为结构化中间数据")
    parser.add_argument("input", type=Path)
    parser.add_argument("--root", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    write_json(args.output, parse_file(args.input, root=args.root))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
